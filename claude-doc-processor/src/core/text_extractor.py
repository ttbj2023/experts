"""
文本提取器

从数字文档直接提取文本内容（跳过OCR）
"""

import os
import fitz  # PyMuPDF
import logging
from typing import Dict, List, Tuple
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph


logger = logging.getLogger(__name__)


class TextExtractor:
    """文本提取器"""

    def __init__(self, config: Dict = None):
        """
        初始化文本提取器

        Args:
            config: 配置字典
        """
        self.config = config or {}
        self.extraction_config = self.config.get('fast_extraction', {})

    # ========================================
    # PDF文本提取
    # ========================================

    def extract_from_digital_pdf(self, pdf_path: str) -> Tuple[str, List[str]]:
        """
        从数字版PDF提取文本和图片

        Args:
            pdf_path: PDF文件路径

        Returns:
            (Markdown文本, 图片路径列表)
        """
        logger.info(f"开始从数字版PDF提取文本: {pdf_path}")

        try:
            doc = fitz.open(pdf_path)

            markdown_content = []
            image_paths = []

            for page_num, page in enumerate(doc):
                # 提取文本块
                blocks = page.get_text("dict")["blocks"]

                page_text = []
                for block in blocks:
                    if block["type"] == 0:  # 文本块
                        for line in block["lines"]:
                            line_text = " ".join([span["text"] for span in line["spans"]])
                            if line_text.strip():
                                page_text.append(line_text)

                # 添加页面内容
                if page_text:
                    markdown_content.append("\n".join(page_text))
                    markdown_content.append("\n\n")

                # 提取图片（如果配置了）
                if self.extraction_config.get('pdf', {}).get('extract_images', True):
                    page_images = self._extract_images_from_page(page, page_num, pdf_path)
                    image_paths.extend(page_images)

            doc.close()

            # 合并所有文本
            full_text = "".join(markdown_content).strip()

            logger.info(f"PDF文本提取完成: 文本长度={len(full_text)}, 图片数={len(image_paths)}")

            return (full_text, image_paths)

        except Exception as e:
            logger.error(f"PDF文本提取失败: {e}")
            raise

    def _extract_images_from_page(
        self,
        page: fitz.Page,
        page_num: int,
        pdf_path: str
    ) -> List[str]:
        """
        从PDF页面提取图片

        Args:
            page: PyMuPDF页面对象
            page_num: 页码
            pdf_path: PDF文件路径

        Returns:
            图片路径列表
        """
        image_list = page.get_images()
        image_paths = []

        if not image_list:
            return image_paths

        # 创建输出目录
        base_name = os.path.splitext(os.path.basename(pdf_path))[0]
        output_dir = os.path.join("output", base_name, "extracted_images")
        os.makedirs(output_dir, exist_ok=True)

        for img_index, img in enumerate(image_list, 1):
            try:
                # 提取图片
                xref = img[0]
                base_image = page.parent.extract_image(xref)

                # 保存图片
                image_ext = base_image["ext"]
                image_data = base_image["image"]

                image_filename = f"page_{page_num + 1:03d}_img_{img_index:02d}.{image_ext}"
                image_path = os.path.join(output_dir, image_filename)

                with open(image_path, "wb") as f:
                    f.write(image_data)

                image_paths.append(image_path)
                logger.debug(f"提取图片: {image_path}")

            except Exception as e:
                logger.warning(f"图片提取失败 (page={page_num + 1}, img={img_index}): {e}")

        return image_paths

    # ========================================
    # DOCX文本提取
    # ========================================

    def extract_from_simple_docx(self, docx_path: str) -> Tuple[str, List[str]]:
        """
        从简单DOCX提取文本（保留结构）

        Args:
            docx_path: DOCX文件路径

        Returns:
            (Markdown文本, 图片路径列表)
        """
        logger.info(f"开始从简单DOCX提取文本: {docx_path}")

        try:
            doc = DocxDocument(docx_path)

            markdown_content = []
            image_paths = []

            # 配置选项
            preserve_tables = self.extraction_config.get('docx', {}).get('preserve_tables', True)
            preserve_lists = self.extraction_config.get('docx', {}).get('preserve_lists', True)

            # 提取段落和表格
            i = 0
            while i < len(doc.element.body):
                element = doc.element.body[i]

                # 处理段落
                if element.tag.endswith('p'):
                    para = None
                    for p in doc.paragraphs:
                        if p._element == element:
                            para = p
                            break

                    if para:
                        para_text = self._process_paragraph(para, preserve_lists)
                        if para_text:
                            markdown_content.append(para_text)
                            markdown_content.append("\n")

                # 处理表格
                elif element.tag.endswith('tbl'):
                    table = None
                    for t in doc.tables:
                        if t._element == element:
                            table = t
                            break

                    if table and preserve_tables:
                        table_text = self._process_table(table)
                        if table_text:
                            markdown_content.append(table_text)
                            markdown_content.append("\n\n")

                i += 1

            # 提取图片（如果有的话）
            image_paths = self._extract_images_from_docx(doc, docx_path)

            # 合并所有文本
            full_text = "".join(markdown_content).strip()

            logger.info(f"DOCX文本提取完成: 文本长度={len(full_text)}, 图片数={len(image_paths)}")

            return (full_text, image_paths)

        except Exception as e:
            logger.error(f"DOCX文本提取失败: {e}")
            raise

    def _process_paragraph(self, para: Paragraph, preserve_lists: bool) -> str:
        """
        处理单个段落，转换为Markdown格式

        Args:
            para: python-docx段落对象
            preserve_lists: 是否保留列表结构

        Returns:
            Markdown文本
        """
        text = para.text.strip()

        if not text:
            return ""

        # 检测标题
        if para.style.name.startswith('Heading'):
            level = para.style.name.replace('Heading ', '')
            try:
                level_num = int(level)
                return "#" * level_num + " " + text + "\n"
            except ValueError:
                pass

        # 检测列表（如果启用）
        if preserve_lists:
            # 检测编号列表
            if 'List Number' in para.style.name:
                return "1. " + text

            # 检测项目符号
            if 'List Bullet' in para.style.name:
                return "- " + text

        # 普通段落
        return text + "\n"

    def _process_table(self, table: Table) -> str:
        """
        处理表格，转换为Markdown表格格式

        Args:
            table: python-docx表格对象

        Returns:
            Markdown表格文本
        """
        if not table.rows:
            return ""

        markdown_rows = []

        # 处理表头
        header_cells = [cell.text.strip() for cell in table.rows[0].cells]
        markdown_rows.append("| " + " | ".join(header_cells) + " |")
        markdown_rows.append("|" + "|".join(["---"] * len(header_cells)) + "|")

        # 处理数据行
        for row in table.rows[1:]:
            cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            markdown_rows.append("| " + " | ".join(cells) + " |")

        return "\n".join(markdown_rows) + "\n"

    def _extract_images_from_docx(
        self,
        doc: DocxDocument,
        docx_path: str
    ) -> List[str]:
        """
        从DOCX提取图片

        Args:
            doc: python-docx文档对象
            docx_path: DOCX文件路径

        Returns:
            图片路径列表
        """
        image_paths = []

        # 创建输出目录
        base_name = os.path.splitext(os.path.basename(docx_path))[0]
        output_dir = os.path.join("output", base_name, "extracted_images")
        os.makedirs(output_dir, exist_ok=True)

        # 遍历所有关系（图片）
        for rel_id, rel in doc.part.rels.items():
            if "image" not in rel.target_ref:
                continue

            try:
                # 读取图片数据
                image_data = rel.target_part.blob

                # 获取图片扩展名
                image_ext = rel.target_ref.split('.')[-1]
                if image_ext not in ['png', 'jpg', 'jpeg', 'gif', 'bmp']:
                    image_ext = 'png'

                # 保存图片
                image_filename = f"image_{rel_id}.{image_ext}"
                image_path = os.path.join(output_dir, image_filename)

                with open(image_path, "wb") as f:
                    f.write(image_data)

                image_paths.append(image_path)
                logger.debug(f"提取图片: {image_path}")

            except Exception as e:
                logger.warning(f"图片提取失败 (rel_id={rel_id}): {e}")

        return image_paths

    # ========================================
    # 格式化Markdown
    # ========================================

    def clean_extracted_text(self, text: str) -> str:
        """
        清理提取的文本

        Args:
            text: 原始文本

        Returns:
            清理后的文本
        """
        # 移除多余空行
        import re
        text = re.sub(r'\n{3,}', '\n\n', text)

        # 移除行首行尾空格
        lines = [line.strip() for line in text.split('\n')]
        text = '\n'.join(lines)

        return text.strip()
